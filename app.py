import json
from pathlib import Path
from html import escape

import requests
import streamlit as st
from docx import Document


# ------------------------------------------------------------------
# Configuration
# ------------------------------------------------------------------

API_BASE_URL = "http://127.0.0.1:8000"

PROJECT_ROOT = Path(__file__).resolve().parent

CASE_INFORMATION = (
    PROJECT_ROOT
    / "data"
    / "input"
    / "03_Case_Information.pdf"
)

GENERATED_AFFIDAVIT = (
    PROJECT_ROOT
    / "outputs"
    / "affidavit_in_reply.docx"
)

GENERATED_REPORT = (
    PROJECT_ROOT
    / "outputs"
    / "evaluation_report.json"
)


# ------------------------------------------------------------------
# Page configuration
# ------------------------------------------------------------------

st.set_page_config(
    page_title="Legal Document Generation & Evaluation Agent",
    page_icon="⚖️",
    layout="wide",
)


# ------------------------------------------------------------------
# Application styling
# ------------------------------------------------------------------

st.markdown(
    """
    <style>

    /* ============================================================
       Global typography
       ============================================================ */

    html,
    body,
    [class*="css"],
    .stApp {
        font-family:
            -apple-system,
            BlinkMacSystemFont,
            "Segoe UI",
            Roboto,
            Helvetica,
            Arial,
            sans-serif;
    }

    .block-container {
        padding-top: 2.5rem;
        padding-bottom: 3rem;
        max-width: 1400px;
    }

    /* ============================================================
       Headings
       ============================================================ */

    h1 {
        font-size: 2.5rem !important;
        font-weight: 700 !important;
        letter-spacing: -0.025em;
    }

    h2 {
        font-size: 1.7rem !important;
        font-weight: 650 !important;
    }

    h3 {
        font-size: 1.25rem !important;
        font-weight: 650 !important;
    }

    /* ============================================================
       Section labels
       ============================================================ */

    .section-label {
        font-size: 0.78rem;
        font-weight: 650;
        text-transform: uppercase;
        letter-spacing: 0.08em;
        opacity: 0.62;
        margin-bottom: 0.35rem;
    }

    .page-subtitle {
        font-size: 1rem;
        opacity: 0.68;
        margin-top: -0.6rem;
        margin-bottom: 1.7rem;
    }

    /* ============================================================
       Document viewer
       ============================================================ */

    .document-heading {
        font-size: 1.45rem;
        font-weight: 650;
        margin-bottom: 0.2rem;
    }

    .document-description {
        font-size: 0.92rem;
        opacity: 0.65;
        margin-bottom: 1rem;
    }

    /* ============================================================
       Generated legal document preview
       ============================================================ */

    .legal-document {
        background: #ffffff;
        color: #111111;
        border-radius: 6px;
        padding: 60px 72px;
        margin-top: 0.5rem;
        margin-bottom: 1rem;
        box-shadow:
            0 2px 10px rgba(0, 0, 0, 0.18);
        font-family:
            "Times New Roman",
            Times,
            serif;
        font-size: 12pt;
        line-height: 1.55;
    }

    .legal-paragraph {
        margin: 0 0 16px 0;
        white-space: pre-wrap;
    }

    .legal-centered {
        text-align: center;
        font-weight: bold;
        margin: 0 0 18px 0;
    }

    .legal-cause-title {
        margin: 22px 0;
        line-height: 1.6;
    }

    .legal-numbered {
        margin: 0 0 16px 0;
        line-height: 1.55;
    }

    .legal-small {
        font-size: 11pt;
    }

    /* ============================================================
       Evaluation
       ============================================================ */

    .evaluation-note {
        font-size: 0.9rem;
        opacity: 0.68;
        margin-bottom: 0.8rem;
    }

    /* ============================================================
       Buttons
       ============================================================ */

    .stButton > button,
    .stDownloadButton > button {
        font-family: inherit;
        font-weight: 600;
    }

    /* ============================================================
       Metrics
       ============================================================ */

    [data-testid="stMetricLabel"] {
        font-size: 0.84rem;
    }

    [data-testid="stMetricValue"] {
        font-size: 1.75rem;
    }

    </style>
    """,
    unsafe_allow_html=True,
)


# ------------------------------------------------------------------
# Helper functions
# ------------------------------------------------------------------

def display_pdf(pdf_path: Path) -> None:
    """
    Display the original case-information PDF using Streamlit's
    native PDF viewer.

    The original PDF is preserved rather than reconstructed as text.
    """

    if not pdf_path.exists():
        st.error(
            f"Document not found: {pdf_path.name}"
        )
        return

    try:
        pdf_bytes = pdf_path.read_bytes()

        st.pdf(
            pdf_bytes,
            height=900,
        )

    except Exception as exc:
        st.error(
            f"Unable to display {pdf_path.name}: {exc}"
        )


def get_docx_blocks() -> list[dict]:
    """
    Read the generated DOCX and return its paragraphs/tables in a
    simple representation suitable for the in-app preview.
    """

    if not GENERATED_AFFIDAVIT.exists():
        return []

    try:
        document = Document(
            GENERATED_AFFIDAVIT
        )

        blocks = []

        # ----------------------------------------------------------
        # Paragraphs
        # ----------------------------------------------------------

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if not text:
                continue

            alignment = paragraph.alignment

            if alignment is not None:
                try:
                    alignment_value = alignment.name
                except Exception:
                    alignment_value = None
            else:
                alignment_value = None

            blocks.append(
                {
                    "type": "paragraph",
                    "text": text,
                    "alignment": alignment_value,
                }
            )

        # ----------------------------------------------------------
        # Tables
        # ----------------------------------------------------------

        for table in document.tables:

            rows = []

            for row in table.rows:

                cells = []

                for cell in row.cells:

                    cell_text = "\n".join(
                        paragraph.text.strip()
                        for paragraph in cell.paragraphs
                        if paragraph.text.strip()
                    )

                    cells.append(
                        cell_text
                    )

                rows.append(cells)

            blocks.append(
                {
                    "type": "table",
                    "rows": rows,
                }
            )

        return blocks

    except Exception:
        return []


def render_generated_document_preview() -> None:
    """
    Render the generated affidavit as a readable legal-document-style
    preview inside the Streamlit application.
    """

    blocks = get_docx_blocks()

    if not blocks:
        st.warning(
            "The generated affidavit exists, but a preview could "
            "not be created."
        )
        return

    html_parts = [
        '<div class="legal-document">'
    ]

    for block in blocks:

        if block["type"] == "paragraph":

            text = escape(
                block["text"]
            )

            alignment = block.get(
                "alignment"
            )

            # ------------------------------------------------------
            # Centered headings
            # ------------------------------------------------------

            if alignment == "CENTER":

                html_parts.append(
                    '<div class="legal-centered">'
                    f"{text}"
                    "</div>"
                )

            # ------------------------------------------------------
            # Regular legal paragraphs
            # ------------------------------------------------------

            else:

                html_parts.append(
                    '<div class="legal-paragraph">'
                    f"{text}"
                    "</div>"
                )

        elif block["type"] == "table":

            html_parts.append(
                '<div class="legal-cause-title">'
            )

            for row in block["rows"]:

                row_text = " ".join(
                    escape(cell)
                    for cell in row
                    if cell.strip()
                )

                if row_text:

                    html_parts.append(
                        f'<div class="legal-paragraph">'
                        f"{row_text}"
                        "</div>"
                    )

            html_parts.append(
                "</div>"
            )

    html_parts.append(
        "</div>"
    )

    st.markdown(
        "".join(html_parts),
        unsafe_allow_html=True,
    )


def get_generated_affidavit_bytes() -> bytes | None:
    """
    Retrieve the generated DOCX through the FastAPI download endpoint.
    """

    try:

        response = requests.get(
            f"{API_BASE_URL}/download/affidavit",
            timeout=30,
        )

        if response.ok:
            return response.content

    except requests.RequestException:
        pass

    # Fallback to the local generated artifact.

    if GENERATED_AFFIDAVIT.exists():

        try:
            return GENERATED_AFFIDAVIT.read_bytes()
        except OSError:
            pass

    return None


def render_download_button() -> None:
    """
    Render the single download button for the generated affidavit.
    """

    document_bytes = (
        get_generated_affidavit_bytes()
    )

    if document_bytes is None:

        st.warning(
            "The generated affidavit is available locally, "
            "but the download service could not be reached."
        )

        return

    st.download_button(
        label="Download Generated Affidavit (.docx)",
        data=document_bytes,
        file_name="affidavit_in_reply.docx",
        mime=(
            "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.document"
        ),
        use_container_width=True,
        key="download_generated_affidavit",
    )


def generate_affidavit() -> dict | None:
    """
    Run the generation workflow through the FastAPI backend.

    When the API is unavailable, which is expected in the public
    Streamlit Cloud deployment because the production LLM service
    is local to the development machine, fall back to the verified
    pre-generated assignment artifacts stored in /outputs.
    """

    try:

        response = requests.post(
            f"{API_BASE_URL}/generate",
            timeout=600,
        )

        if response.ok:
            return response.json()

        try:

            detail = response.json().get(
                "detail",
                response.text,
            )

        except Exception:

            detail = response.text

        st.warning(
            f"Generation API unavailable "
            f"(HTTP {response.status_code}). "
            "Showing the verified pre-generated demo output."
        )

    except requests.RequestException:

        st.info(
            "Demo Mode: the local generation API is not available "
            "in the cloud deployment. Showing the verified "
            "pre-generated assignment output."
        )

    # ----------------------------------------------------------
    # Cloud demo fallback
    # ----------------------------------------------------------

    if not GENERATED_REPORT.exists():

        st.error(
            "No generation API is available and no pre-generated "
            "evaluation report was found."
        )

        return None

    try:

        report = json.loads(
            GENERATED_REPORT.read_text(
                encoding="utf-8"
            )
        )

    except (OSError, ValueError, TypeError) as exc:

        st.error(
            f"Unable to load the pre-generated demo output: {exc}"
        )

        return None

    return {
        "status": "demo_mode",
        "revision_count": report.get(
            "revision_count",
            0,
        ),
        "evaluation_history_count": report.get(
            "evaluation_history_count",
            1,
        ),
        "report": report,
        "artifacts": {
            "affidavit": "/download/affidavit",
            "evaluation_report": "/report",
        },
    }


# ------------------------------------------------------------------
# Header
# ------------------------------------------------------------------

st.title(
    "Legal Document Generation & Evaluation Agent"
)

st.markdown(
    '<div class="page-subtitle">'
    'Generate, validate and evaluate an Affidavit in Reply '
    'from the supplied case information.'
    '</div>',
    unsafe_allow_html=True,
)


# ------------------------------------------------------------------
# Documents
# ------------------------------------------------------------------

st.markdown(
    '<div class="section-label">Documents</div>',
    unsafe_allow_html=True,
)

document_col1, document_col2 = st.columns(2)


if "active_document" not in st.session_state:

    st.session_state[
        "active_document"
    ] = "Original Case Information"


with document_col1:

    if st.button(
        "Original Case Information",
        use_container_width=True,
    ):

        st.session_state[
            "active_document"
        ] = "Original Case Information"

        st.rerun()


with document_col2:

    if st.button(
        "Generated Affidavit",
        use_container_width=True,
    ):

        st.session_state[
            "active_document"
        ] = "Generated Affidavit"

        st.rerun()


st.divider()


# ------------------------------------------------------------------
# Original Case Information
# ------------------------------------------------------------------

active_document = st.session_state[
    "active_document"
]


if active_document == "Original Case Information":

    st.markdown(
        '<div class="document-heading">'
        'Original Case Information'
        '</div>',
        unsafe_allow_html=True,
    )

    st.markdown(
        '<div class="document-description">'
        'The original case-information PDF supplied for the assignment. '
        'This document is used as the source for case and entity extraction.'
        '</div>',
        unsafe_allow_html=True,
    )

    display_pdf(
        CASE_INFORMATION
    )


# ------------------------------------------------------------------
# Generated Affidavit
# ------------------------------------------------------------------

elif active_document == "Generated Affidavit":

    st.markdown(
        '<div class="document-heading">'
        'Generated Affidavit in Reply'
        '</div>',
        unsafe_allow_html=True,
    )

    st.markdown(
        '<div class="document-description">'
        'The latest affidavit produced by the document-generation workflow, '
        'shown in a legal-document-style preview.'
        '</div>',
        unsafe_allow_html=True,
    )

    if GENERATED_AFFIDAVIT.exists():

        render_generated_document_preview()

        render_download_button()

    else:

        st.info(
            "No generated affidavit is available yet. "
            "Run the generation workflow below."
        )


# ------------------------------------------------------------------
# Generation
# ------------------------------------------------------------------

st.divider()

st.markdown(
    '<div class="section-label">Generation</div>',
    unsafe_allow_html=True,
)

st.subheader(
    "Generate Affidavit in Reply"
)

st.write(
    "Run the complete extraction, mapping, drafting, "
    "validation and evaluation workflow."
)

if st.button(
    "Generate Affidavit in Reply",
    type="primary",
    use_container_width=True,
):

    with st.spinner(
        "Generating and evaluating the Affidavit in Reply..."
    ):

        result = generate_affidavit()

    if result:

        st.session_state[
            "generation_result"
        ] = result

        st.session_state[
            "active_document"
        ] = "Generated Affidavit"

        if result.get("status") == "demo_mode":

            st.success(
                "Verified demo affidavit and evaluation report loaded successfully."
            )

        else:

            st.success(
                "Affidavit generated and evaluated successfully."
            )

        st.rerun()


# ------------------------------------------------------------------
# Evaluation
# ------------------------------------------------------------------

if "generation_result" in st.session_state:

    result = st.session_state[
        "generation_result"
    ]

    report = result[
        "report"
    ]

    st.divider()

    st.markdown(
        '<div class="section-label">Evaluation</div>',
        unsafe_allow_html=True,
    )

    st.subheader(
        "Latest Evaluation"
    )

    # --------------------------------------------------------------
    # Overall result
    # --------------------------------------------------------------

    summary_col1, summary_col2, summary_col3 = (
        st.columns(3)
    )

    with summary_col1:

        st.metric(
            "Overall Score",
            f'{report["overall_score"]:.1f} / 100',
        )

    with summary_col2:

        st.metric(
            "Status",
            report["status"],
        )

    with summary_col3:

        st.metric(
            "Issues",
            len(report["issues"]),
        )


    # --------------------------------------------------------------
    # Six required assignment dimensions
    # --------------------------------------------------------------

    st.markdown(
        "### Evaluation Dimensions"
    )

    st.markdown(
        '<div class="evaluation-note">'
        'Six required dimensions used to evaluate the generated affidavit.'
        '</div>',
        unsafe_allow_html=True,
    )

    dimension_scores = report[
        "dimension_scores"
    ]

    dimensions = [
        "Entity Accuracy",
        "Completeness",
        "Structure",
        "Consistency",
        "Template Fidelity",
        "Hallucination Check",
    ]

    dimension_row_1 = st.columns(3)

    for column, dimension in zip(
        dimension_row_1,
        dimensions[:3],
    ):

        with column:

            st.metric(
                dimension,
                f'{dimension_scores[dimension]:.1f}',
            )


    dimension_row_2 = st.columns(3)

    for column, dimension in zip(
        dimension_row_2,
        dimensions[3:],
    ):

        with column:

            st.metric(
                dimension,
                f'{dimension_scores[dimension]:.1f}',
            )


    # --------------------------------------------------------------
    # Validation summary
    # --------------------------------------------------------------

    st.markdown(
        "### Validation"
    )

    validation_col1, validation_col2 = (
        st.columns(2)
    )

    deterministic = (
        report[
            "validation_summary"
        ][
            "deterministic"
        ]
    )

    semantic = (
        report[
            "validation_summary"
        ][
            "semantic"
        ]
    )

    with validation_col1:

        st.markdown(
            "**Deterministic Validation**"
        )

        if deterministic["passed"]:

            st.success(
                f'Passed — {deterministic["score"]:.1f}/100'
            )

        else:

            st.error(
                f'Failed — '
                f'{deterministic["error_count"]} error(s)'
            )


    with validation_col2:

        st.markdown(
            "**Semantic Evaluation**"
        )

        if semantic["passed"]:

            st.success(
                f'Passed — '
                f'{semantic["issue_count"]} issue(s)'
            )

        else:

            st.error(
                f'Failed — '
                f'{semantic["issue_count"]} issue(s)'
            )


    # --------------------------------------------------------------
    # Scoring methodology
    # --------------------------------------------------------------

    with st.expander(
        "Scoring methodology and weights"
    ):

        st.write(
            report[
                "scoring"
            ][
                "method"
            ]
        )

        st.write(
            "**Evaluation weights**"
        )

        weights = report[
            "scoring"
        ][
            "weights"
        ]

        for dimension in dimensions:

            st.write(
                f"- **{dimension}:** "
                f"{weights[dimension]:.0%}"
            )


    # --------------------------------------------------------------
    # Issues
    # --------------------------------------------------------------

    st.markdown(
        "### Detected Issues"
    )

    if report["issues"]:

        for issue in report["issues"]:

            st.warning(
                issue
            )

    else:

        st.success(
            "No issues were detected in the latest supplied-case run."
        )


    # --------------------------------------------------------------
    # Semantic observation
    # --------------------------------------------------------------

    semantic_summary = semantic.get(
        "summary",
        "",
    )

    if semantic_summary:

        with st.expander(
            "Semantic evaluator observation"
        ):

            st.write(
                semantic_summary
            )


    # --------------------------------------------------------------
    # Robustness testing
    # --------------------------------------------------------------

    st.divider()

    st.markdown(
        '<div class="section-label">Robustness Testing</div>',
        unsafe_allow_html=True,
    )

    st.subheader(
        "Adversarial Validation Evidence"
    )

    st.write(
        "Separate stress-test results from deliberately mutated "
        "document outputs. These figures are not substituted for "
        "the latest production evaluation score."
    )

    robustness_col1, robustness_col2 = (
        st.columns(2)
    )

    with robustness_col1:

        st.markdown(
            "#### Deterministic Adversarial Audit"
        )

        st.metric(
            "Detection Rate",
            "89.5%",
        )

        st.write(
            "**19** mutations tested"
        )

        st.write(
            "**17** caught · **2** missed"
        )


    with robustness_col2:

        st.markdown(
            "#### Semantic Adversarial Audit"
        )

        st.metric(
            "Detection Rate",
            "50.0%",
        )

        st.write(
            "**4** hallucination mutations tested"
        )

        st.write(
            "**2** caught · **2** missed"
        )


    # --------------------------------------------------------------
    # Workflow metadata
    # --------------------------------------------------------------

    st.markdown(
        "### Workflow Metadata"
    )

    metadata_col1, metadata_col2, metadata_col3 = (
        st.columns(3)
    )

    with metadata_col1:

        st.metric(
            "Revision Count",
            result.get(
                "revision_count",
                0,
            ),
        )

    with metadata_col2:

        st.metric(
            "Evaluation History",
            result.get(
                "evaluation_history_count",
                0,
            ),
        )

    with metadata_col3:

        st.metric(
            "Evaluation Model",
            report.get(
                "model",
                "N/A",
            ),
        )


    # --------------------------------------------------------------
    # Artifact controls
    # --------------------------------------------------------------

    st.markdown(
        "### Generated Artifact"
    )

    if st.button(
        "Open Generated Affidavit",
        use_container_width=True,
    ):

        st.session_state[
            "active_document"
        ] = "Generated Affidavit"

        st.rerun()