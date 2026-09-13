from pathlib import Path

from docx import Document

from src.output.docx_renderer import render_affidavit_docx


def test_render_affidavit_docx(tmp_path):
    generated_draft = {
        "document_type": "Affidavit in Reply",
        "sections": {
            "FORUM_HEADING": "IN THE HIGH COURT OF JUDICATURE AT BOMBAY",
            "JURISDICTION": "CIVIL APPELLATE JURISDICTION",
            "CASE_NUMBER": "WRIT PETITION NO. 3147 OF 2026",
            "CAUSE_TITLE": (
                "Arjun Mehta, Age 38 years, Occupation: Business, "
                "residing at Mumbai. ...Petitioner\n"
                "VERSUS\n"
                "1. State of Maharashtra, Through the Principal Secretary, "
                "Mumbai. ...Respondent No.1\n"
                "2. Rohan Deshpande, Age 42 years, residing at Mumbai. "
                "...Respondent No.2"
            ),
            "AFFIDAVIT_TITLE": (
                "AFFIDAVIT IN REPLY ON BEHALF OF RESPONDENT NO. 2"
            ),
            "DEPONENT_CLAUSE": (
                "I, Rohan Deshpande, Age 42 years, residing at Mumbai, "
                "the Respondent No.2 above named, do hereby solemnly affirm "
                "and state as under:"
            ),
        },
        "paragraphs": [
            (
                "I say that I am the Respondent No.2 in the above Writ Petition "
                "and am well acquainted with the facts and circumstances of the "
                "case. I have perused the Petition and the documents annexed "
                "thereto and am competent to affirm this Affidavit in Reply."
            ),
            (
                "At the outset, I deny each and every allegation, contention and "
                "submission made in the Writ Petition, save and except those "
                "specifically admitted herein. I say that the Petition is "
                "misconceived, devoid of merits and is liable to be dismissed "
                "in limine."
            ),
            (
                "I say that the Petitioner has suppressed material facts and "
                "prior correspondence having a direct bearing on the controversy. "
                "The action complained of has been taken strictly in accordance "
                "with law and after following due procedure."
            ),
        ],
        "exhibit": "EXHIBIT-‘A’",
        "prayer": {
            "opening": (
                "I therefore respectfully pray that this Hon'ble Court "
                "may be pleased to:"
            ),
            "items": [
                "dismiss the present Writ Petition with costs;",
                "refuse any interim or ad-interim relief sought by the Petitioner;",
                (
                    "grant such other and further reliefs as this Hon'ble Court "
                    "may deem fit and proper in the facts and circumstances "
                    "of the case."
                ),
            ],
        },
        "jurat": {
            "affirmed": "Solemnly affirmed at Mumbai",
            "date": "On this 5th day of September 2026",
            "before_me": "Before Me",
        },
        "verification": {
            "body": (
                "I, Rohan Deshpande, the Deponent above named, do hereby "
                "verify that the contents of paragraphs 1 to 3 and the "
                "Prayer above are true and correct to my knowledge and "
                "belief and that nothing material has been concealed "
                "therefrom."
            ),
            "verified_at": (
                "Verified at Mumbai on this 5th day of September 2026."
            ),
            "deponent": "DEPONENT",
        },
        "advocate": {
            "firm": "MEHTA & KULKARNI",
            "acting_for": "Advocates for the Respondent No.2.",
        },
    }

    output_path = tmp_path / "affidavit_in_reply.docx"

    # ---------------------------------------------------------
    # Render
    # ---------------------------------------------------------
    result = render_affidavit_docx(
        generated_draft,
        output_path,
    )

    # ---------------------------------------------------------
    # File-level checks
    # ---------------------------------------------------------
    assert result == output_path
    assert output_path.exists()
    assert output_path.stat().st_size > 0

    # ---------------------------------------------------------
    # Open generated DOCX
    # ---------------------------------------------------------
    document = Document(output_path)

    # ---------------------------------------------------------
    # Collect normal paragraph text
    # ---------------------------------------------------------
    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    # ---------------------------------------------------------
    # Collect text contained inside tables
    #
    # The renderer intentionally uses tables for:
    # - cause-title left/right party layout
    # - jurat "Before Me" / "DEPONENT" layout
    #
    # python-docx does not include table paragraphs in
    # document.paragraphs, so they must be collected separately.
    # ---------------------------------------------------------
    table_text = []

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    text = paragraph.text.strip()

                    if text:
                        table_text.append(text)

    # Combine document-level and table-level text.
    full_text = "\n".join(paragraphs + table_text)

    # ---------------------------------------------------------
    # Core document sections
    # ---------------------------------------------------------
    assert "IN THE HIGH COURT OF JUDICATURE AT BOMBAY" in full_text
    assert "CIVIL APPELLATE JURISDICTION" in full_text
    assert "WRIT PETITION NO. 3147 OF 2026" in full_text
    assert "VERSUS" in full_text
    assert (
        "AFFIDAVIT IN REPLY ON BEHALF OF RESPONDENT NO. 2"
        in full_text
    )

    # ---------------------------------------------------------
    # Cause title
    # ---------------------------------------------------------
    assert "Arjun Mehta" in full_text
    assert "...Petitioner" in full_text

    assert "State of Maharashtra" in full_text
    assert "...Respondent No.1" in full_text

    assert "Rohan Deshpande" in full_text
    assert "...Respondent No.2" in full_text

    # ---------------------------------------------------------
    # Deponent clause
    # ---------------------------------------------------------
    assert "Rohan Deshpande" in full_text
    assert (
        "the Respondent No.2 above named"
        in full_text
    )
    assert (
        "do hereby solemnly affirm and state as under:"
        in full_text
    )

    # ---------------------------------------------------------
    # Numbered body paragraphs
    # ---------------------------------------------------------
    assert any(
        paragraph.startswith("1.")
        for paragraph in paragraphs
    )

    assert any(
        paragraph.startswith("2.")
        for paragraph in paragraphs
    )

    assert any(
        paragraph.startswith("3.")
        for paragraph in paragraphs
    )

    # ---------------------------------------------------------
    # Body paragraph content
    # ---------------------------------------------------------
    assert (
        "I say that I am the Respondent No.2"
        in full_text
    )

    assert (
        "At the outset, I deny each and every allegation"
        in full_text
    )

    assert (
        "The action complained of has been taken strictly "
        "in accordance with law"
        in full_text
    )

    # ---------------------------------------------------------
    # Exhibit
    # ---------------------------------------------------------
    assert "EXHIBIT-‘A’" in full_text

    # ---------------------------------------------------------
    # Prayer
    # ---------------------------------------------------------
    assert "PRAYER" in full_text

    assert (
        "(a) dismiss the present Writ Petition with costs;"
        in full_text
    )

    assert (
        "(b) refuse any interim or ad-interim relief sought by "
        "the Petitioner;"
        in full_text
    )

    assert (
        "(c) grant such other and further reliefs as this Hon'ble "
        "Court may deem fit and proper in the facts and circumstances "
        "of the case."
        in full_text
    )

    # ---------------------------------------------------------
    # Jurat
    # ---------------------------------------------------------
    assert "Solemnly affirmed at Mumbai" in full_text

    assert (
        "On this 5th day of September 2026"
        in full_text
    )

    assert "Before Me" in full_text
    assert "DEPONENT" in full_text

    # ---------------------------------------------------------
    # Verification
    # ---------------------------------------------------------
    assert "VERIFICATION" in full_text

    assert (
        "paragraphs 1 to 3"
        in full_text
    )

    assert (
        "Prayer above are true and correct to my knowledge and belief"
        in full_text
    )

    assert (
        "nothing material has been concealed therefrom"
        in full_text
    )

    assert (
        "Verified at Mumbai on this 5th day of September 2026."
        in full_text
    )

    # ---------------------------------------------------------
    # Advocate / drafting block
    # ---------------------------------------------------------
    assert "MEHTA & KULKARNI" in full_text

    assert (
        "Advocates for the Respondent No.2."
        in full_text
    )